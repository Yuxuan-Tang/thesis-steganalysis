#!/usr/bin/env python
"""Build the GitHub open-source platform submission note for the college."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "GitHub开源平台提交说明.docx"


def set_run_font(run, name: str = "Microsoft YaHei", size: int | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def style_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for style_name, size, color, before, after in (
        ("Title", 20, "0B2545", 0, 8),
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
    ):
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_kv_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    widths = (Inches(1.6), Inches(4.9))
    for row_idx, (key, value) in enumerate(rows):
        cells = table.rows[row_idx].cells
        for idx, width in enumerate(widths):
            cells[idx].width = width
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[idx])
        set_cell_shading(cells[0], "F2F4F7")
        cells[0].paragraphs[0].add_run(key)
        cells[1].paragraphs[0].add_run(value)
        for cell in cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    set_run_font(run)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run_font(run)


def main() -> None:
    doc = Document()
    style_doc(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("GitHub 开源平台提交说明")
    set_run_font(run, size=20)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("本科毕业设计代码与必要文档开源说明")
    run.font.color.rgb = RGBColor.from_string("555555")
    set_run_font(run, size=11)

    doc.add_heading("一、项目信息", level=1)
    add_kv_table(
        doc,
        [
            ("项目名称", "面向抗隐写分析的生成式无载体图像隐写安全性评估研究"),
            ("仓库名称", "thesis-steganalysis"),
            ("GitHub 链接", "https://github.com/<你的用户名>/thesis-steganalysis"),
            ("开源许可证", "MIT License（仅适用于本人自研脚本与整理文档）"),
        ],
    )

    doc.add_heading("二、开源内容说明", level=1)
    add_bullets(
        doc,
        [
            "本人编写的实验运行脚本、传统 LSB 对照基线和轻量化隐写分析脚本。",
            "实验指标计算、结果汇总、参数敏感性分析和图表生成相关代码。",
            "少量代表性结果样例、汇总 CSV、实验运行说明和论文写作辅助材料。",
            "仓库不包含虚拟环境、缓存、模型权重和大批量原始实验输出。",
        ],
    )

    doc.add_heading("三、第三方代码与模型说明", level=1)
    p = doc.add_paragraph()
    run = p.add_run(
        "DiffStega 为外部公开开源项目，本提交仓库不再分发其源码、预训练模型或权重文件。"
        "如需复现实验，应由使用者根据 DiffStega 原项目说明自行获取第三方代码和 IP-Adapter 模型，"
        "并遵守对应项目许可证、模型协议和平台使用条款。"
    )
    set_run_font(run)

    doc.add_heading("四、运行环境与复现方式", level=1)
    add_bullets(
        doc,
        [
            "主要实验脚本面向 Windows、PowerShell、Conda、Python 3.11 和 CUDA 12.x 环境。",
            "运行 setup_diffstega_env.ps1 可创建本地实验环境并安装 DiffStega 及分析脚本所需依赖。",
            "运行 run_analysis.ps1 可生成 LSB 对照实验与基础指标分析结果。",
            "运行 run_hq_*.ps1 可执行高质量样例、参数校准和扩展结果分析。",
            "仓库 README 已给出目录结构、第三方依赖获取方式和主要脚本用途。",
        ],
    )

    doc.add_heading("五、注意事项", level=1)
    add_bullets(
        doc,
        [
            "GitHub 仓库链接中的 <你的用户名> 需要在正式提交前替换为实际 GitHub 用户名。",
            "如果学院后续提供固定模板，应以学院模板为准替换本说明文档的版式。",
            "本仓库开源范围为毕业设计配套代码和说明材料，不代表对第三方项目重新授权。",
        ],
    )

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("GitHub 开源平台提交说明")
    run.font.color.rgb = RGBColor.from_string("555555")
    set_run_font(run, size=9)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
